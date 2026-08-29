"""
Udyam Verification PDF Report
================================
Reads the JSON output from udyam_autofill.py and generates a
clean professional PDF with:
  - Cover page + summary stats
  - One section per enterprise with all fields
  - Classification history table per enterprise
  - Overall classification summary at the end
"""

import json, os, sys
from datetime import datetime

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    PageBreak, HRFlowable, KeepTogether
)

# ── Palette ───────────────────────────────────────────────
C_BLUE        = colors.HexColor("#1A56DB")
C_BLUE_DARK   = colors.HexColor("#1347C0")
C_BLUE_LIGHT  = colors.HexColor("#EFF6FF")
C_BLUE_MID    = colors.HexColor("#DBEAFE")
C_GREEN       = colors.HexColor("#057A55")
C_GREEN_BG    = colors.HexColor("#D1FAE5")
C_RED         = colors.HexColor("#E02424")
C_RED_BG      = colors.HexColor("#FEE2E2")
C_YELLOW      = colors.HexColor("#C27803")
C_YELLOW_BG   = colors.HexColor("#FEF3C7")
C_GRAY        = colors.HexColor("#6B7280")
C_GRAY_LIGHT  = colors.HexColor("#F9FAFB")
C_GRAY_LINE   = colors.HexColor("#E5E7EB")
C_WHITE       = colors.white
C_BLACK       = colors.HexColor("#111827")

TYPE_COLOR = {
    "Micro":  (colors.HexColor("#1D4ED8"), colors.HexColor("#DBEAFE")),
    "Small":  (colors.HexColor("#065F46"), colors.HexColor("#D1FAE5")),
    "Medium": (colors.HexColor("#92400E"), colors.HexColor("#FEF3C7")),
}
STATUS_COLOR = {
    "Verified":       (C_GREEN,  C_GREEN_BG),
    "Invalid":        (C_RED,    C_RED_BG),
    "Invalid Format": (C_RED,    C_RED_BG),
    "Error":          (C_YELLOW, C_YELLOW_BG),
    "Timeout":        (C_YELLOW, C_YELLOW_BG),
}

W, H = A4


# ── Styles ────────────────────────────────────────────────

def make_styles():
    S = {}

    def ps(name, **kw):
        S[name] = ParagraphStyle(name, **kw)

    ps("cover_title",    fontName="Helvetica-Bold",   fontSize=24, textColor=C_WHITE,      alignment=TA_CENTER, spaceAfter=4)
    ps("cover_sub",      fontName="Helvetica",         fontSize=12, textColor=colors.HexColor("#BFD3F6"), alignment=TA_CENTER, spaceAfter=3)
    ps("cover_date",     fontName="Helvetica",         fontSize=9,  textColor=colors.HexColor("#93C5FD"), alignment=TA_CENTER)

    ps("section_hdr",    fontName="Helvetica-Bold",   fontSize=13, textColor=C_BLUE,       spaceBefore=8, spaceAfter=4)
    ps("enterprise_hdr", fontName="Helvetica-Bold",   fontSize=11, textColor=C_WHITE,      alignment=TA_LEFT)
    ps("field_label",    fontName="Helvetica-Bold",   fontSize=8.5,textColor=C_GRAY)
    ps("field_value",    fontName="Helvetica",         fontSize=8.5,textColor=C_BLACK)
    ps("udyam_id",       fontName="Courier-Bold",      fontSize=8,  textColor=C_BLUE)
    ps("tbl_hdr",        fontName="Helvetica-Bold",   fontSize=8,  textColor=C_WHITE,      alignment=TA_CENTER)
    ps("tbl_cell",       fontName="Helvetica",         fontSize=8,  textColor=C_BLACK,      alignment=TA_CENTER)
    ps("tbl_cell_l",     fontName="Helvetica",         fontSize=8,  textColor=C_BLACK,      alignment=TA_LEFT)
    ps("note",           fontName="Helvetica-Oblique", fontSize=7.5,textColor=C_GRAY,       spaceAfter=4)
    ps("footer_txt",     fontName="Helvetica",         fontSize=7,  textColor=C_GRAY,       alignment=TA_CENTER)

    return S


# ── Numbered canvas (header + footer) ────────────────────

def make_canvas_class(total_records):
    from reportlab.pdfgen import canvas as CV

    class NC(CV.Canvas):
        def __init__(self, *a, **kw):
            super().__init__(*a, **kw)
            self._saved = []

        def showPage(self):
            self._saved.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._saved)
            for state in self._saved:
                self.__dict__.update(state)
                self._draw_chrome(total)
                super().showPage()
            super().save()

        def _draw_chrome(self, total):
            pg = self._pageNumber
            self.saveState()

            # Header (skip cover)
            if pg > 1:
                self.setFillColor(C_BLUE)
                self.rect(0, H - 16*mm, W, 16*mm, fill=1, stroke=0)
                self.setFont("Helvetica-Bold", 9)
                self.setFillColor(C_WHITE)
                self.drawString(15*mm, H - 10*mm, "Udyam Registration Verification Report")
                self.setFont("Helvetica", 8)
                self.drawRightString(W - 15*mm, H - 10*mm,
                    f"Total Records: {total_records}")

            # Footer
            self.setFillColor(C_GRAY_LINE)
            self.rect(0, 0, W, 9*mm, fill=1, stroke=0)
            self.setFont("Helvetica", 7)
            self.setFillColor(C_GRAY)
            self.drawString(15*mm, 3*mm,
                f"Generated: {datetime.now().strftime('%d %B %Y, %H:%M')}")
            self.drawCentredString(W/2, 3*mm, "CONFIDENTIAL — Internal Use Only")
            self.drawRightString(W - 15*mm, 3*mm, f"Page {pg} of {total}")
            self.restoreState()

    return NC


# ── Cover page ────────────────────────────────────────────

def cover_page(S, records):
    elems = []

    total    = len(records)
    verified = sum(1 for r in records if r.get("vstatus") == "Verified")
    invalid  = sum(1 for r in records if "invalid" in r.get("vstatus","").lower())
    errors   = sum(1 for r in records if r.get("vstatus","") in ("Error","Timeout"))
    rate     = f"{verified/total*100:.1f}%" if total else "0%"

    # Blue banner — each row has explicit padding to avoid merging
    banner = Table(
        [
            [Paragraph("Udyam Registration", S["cover_title"])],
            [Paragraph("Verification Report", S["cover_title"])],
            [Spacer(1, 6*mm)],
            [Paragraph("Ministry of MSME — Bulk Verification Results", S["cover_sub"])],
            [Paragraph(datetime.now().strftime("%d %B %Y"), S["cover_date"])],
        ],
        colWidths=[W - 40*mm],
        rowHeights=[14*mm, 14*mm, 6*mm, 8*mm, 7*mm],
    )
    banner.setStyle(TableStyle([
        ("BACKGROUND",    (0,0), (-1,-1), C_BLUE),
        ("TOPPADDING",    (0,0), (-1,-1), 0),
        ("BOTTOMPADDING", (0,0), (-1,-1), 0),
        ("LEFTPADDING",   (0,0), (-1,-1), 10),
        ("RIGHTPADDING",  (0,0), (-1,-1), 10),
        ("VALIGN",        (0,0), (-1,-1), "MIDDLE"),
        # Extra top space before first title row
        ("TOPPADDING",    (0,0), (0,0),   18),
        # Extra bottom space after date row
        ("BOTTOMPADDING", (0,4), (0,4),   14),
    ]))
    elems.append(banner)
    elems.append(Spacer(1, 14*mm))

    # Stat cards
    stat_data = [
        (str(total),    "Total",        C_BLUE,   C_BLUE_LIGHT),
        (str(verified), "Verified",     C_GREEN,  C_GREEN_BG),
        (str(invalid),  "Invalid",      C_RED,    C_RED_BG),
        (str(errors),   "Errors",       C_YELLOW, C_YELLOW_BG),
        (rate,          "Success Rate", C_BLUE,   C_BLUE_LIGHT),
    ]
    cards = []
    for num, lbl, tc, bg in stat_data:
        ns = ParagraphStyle("_n", fontName="Helvetica-Bold", fontSize=22,
                             textColor=tc, alignment=TA_CENTER)
        ls = ParagraphStyle("_l", fontName="Helvetica", fontSize=8,
                             textColor=C_GRAY, alignment=TA_CENTER, spaceBefore=2)
        card = Table([[Paragraph(num, ns)], [Paragraph(lbl, ls)]],
                      colWidths=[30*mm])
        card.setStyle(TableStyle([
            ("BACKGROUND",   (0,0),(-1,-1), bg),
            ("TOPPADDING",   (0,0),(-1,-1), 10),
            ("BOTTOMPADDING",(0,0),(-1,-1), 10),
            ("LEFTPADDING",  (0,0),(-1,-1), 4),
            ("RIGHTPADDING", (0,0),(-1,-1), 4),
        ]))
        cards.append(card)

    stat_row = Table([cards], colWidths=[32*mm]*5)
    stat_row.setStyle(TableStyle([
        ("ALIGN",  (0,0),(-1,-1),"CENTER"),
        ("VALIGN", (0,0),(-1,-1),"MIDDLE"),
        ("LEFTPADDING",  (0,0),(-1,-1), 2),
        ("RIGHTPADDING", (0,0),(-1,-1), 2),
    ]))
    elems.append(stat_row)
    elems.append(Spacer(1, 10*mm))
    elems.append(HRFlowable(width="100%", thickness=1, color=C_GRAY_LINE))
    elems.append(Spacer(1, 6*mm))

    # Info box
    i9 = ParagraphStyle("_i", fontName="Helvetica", fontSize=9,
                          textColor=C_GRAY, leading=14)
    ih = ParagraphStyle("_ih", fontName="Helvetica-Bold", fontSize=10,
                          textColor=C_BLUE)
    info = Table([
        [Paragraph("Report Details", ih)],
        [Paragraph(f"<b>Records:</b> {total}  |  "
                   f"<b>Verified:</b> {verified}  |  "
                   f"<b>Invalid:</b> {invalid}  |  "
                   f"<b>Errors:</b> {errors}", i9)],
        [Paragraph(f"<b>Source:</b> udyamregistration.gov.in", i9)],
        [Paragraph(f"<b>Generated:</b> {datetime.now().strftime('%A, %d %B %Y at %H:%M')}", i9)],
    ], colWidths=[W - 60*mm])
    info.setStyle(TableStyle([
        ("BACKGROUND",   (0,0),(-1,-1), C_BLUE_LIGHT),
        ("BOX",          (0,0),(-1,-1), 1, C_BLUE),
        ("TOPPADDING",   (0,0),(-1,-1), 6),
        ("BOTTOMPADDING",(0,0),(-1,-1), 6),
        ("LEFTPADDING",  (0,0),(-1,-1), 12),
        ("RIGHTPADDING", (0,0),(-1,-1), 12),
    ]))
    elems.append(info)
    elems.append(PageBreak())
    return elems


# ── Single enterprise block ───────────────────────────────

def enterprise_block(S, rec, idx):
    elems = []

    udyam_id = rec.get("udyam_id", "")
    name     = rec.get("name_of_enterprise", "") or udyam_id
    vstatus  = rec.get("vstatus", "Pending")
    stxt, sbg = STATUS_COLOR.get(vstatus, (C_GRAY, C_GRAY_LIGHT))

    # ── Enterprise header bar ────────────────────────────
    status_p = Paragraph(vstatus, ParagraphStyle(
        "_vs", fontName="Helvetica-Bold", fontSize=8,
        textColor=stxt, alignment=TA_CENTER))
    status_cell = Table([[status_p]], colWidths=[24*mm])
    status_cell.setStyle(TableStyle([
        ("BACKGROUND",   (0,0),(-1,-1), sbg),
        ("TOPPADDING",   (0,0),(-1,-1), 4),
        ("BOTTOMPADDING",(0,0),(-1,-1), 4),
        ("LEFTPADDING",  (0,0),(-1,-1), 4),
        ("RIGHTPADDING", (0,0),(-1,-1), 4),
    ]))

    hdr_row = Table([[
        Paragraph(f"{idx}.  {name}", S["enterprise_hdr"]),
        status_cell,
    ]], colWidths=[W - 64*mm - 24*mm, 28*mm])
    hdr_row.setStyle(TableStyle([
        ("BACKGROUND",   (0,0),(0,0), C_BLUE),
        ("BACKGROUND",   (1,0),(1,0), C_BLUE),
        ("VALIGN",       (0,0),(-1,-1), "MIDDLE"),
        ("TOPPADDING",   (0,0),(-1,-1), 8),
        ("BOTTOMPADDING",(0,0),(-1,-1), 8),
        ("LEFTPADDING",  (0,0),(0,0), 12),
        ("RIGHTPADDING", (1,0),(1,0), 6),
    ]))
    elems.append(hdr_row)

    # Udyam ID sub-line
    id_bar = Table([[Paragraph(udyam_id, S["udyam_id"])]],
                    colWidths=[W - 40*mm])
    id_bar.setStyle(TableStyle([
        ("BACKGROUND",   (0,0),(-1,-1), C_BLUE_MID),
        ("TOPPADDING",   (0,0),(-1,-1), 3),
        ("BOTTOMPADDING",(0,0),(-1,-1), 3),
        ("LEFTPADDING",  (0,0),(-1,-1), 12),
    ]))
    elems.append(id_bar)
    elems.append(Spacer(1, 3*mm))

    # ── Fields grid (2 columns) ──────────────────────────
    fields = [
        ("Name of Enterprise",              rec.get("name_of_enterprise",    "") or "—"),
        ("Date of Incorporation",           rec.get("date_of_incorporation", "") or "—"),
        ("Major Activity",                  rec.get("major_activity",        "") or "—"),
        ("Social Category",                 rec.get("social_category",       "") or "—"),
        ("Date of Commencement",            rec.get("date_of_commencement",  "") or "—"),
        ("Organisation Type",               rec.get("org_type",              "") or "—"),
        ("NIC Code",                        rec.get("nic_code",              "") or "—"),
        ("State",                           rec.get("state",                 "") or "—"),
        ("District",                        rec.get("district",              "") or "—"),
        ("Gender",                          rec.get("gender",                "") or "—"),
    ]

    # Arrange into 2 columns
    col_w   = (W - 40*mm) / 2
    grid    = []
    thin    = colors.HexColor("#E5E7EB")

    for i in range(0, len(fields), 2):
        left  = fields[i]
        right = fields[i+1] if i+1 < len(fields) else ("", "")
        row   = [
            Paragraph(left[0],  S["field_label"]),
            Paragraph(left[1],  S["field_value"]),
            Paragraph(right[0], S["field_label"]),
            Paragraph(right[1], S["field_value"]),
        ]
        grid.append(row)

    field_tbl = Table(grid, colWidths=[col_w*0.38, col_w*0.62,
                                        col_w*0.38, col_w*0.62])
    field_tbl_style = [
        ("VALIGN",       (0,0),(-1,-1), "TOP"),
        ("TOPPADDING",   (0,0),(-1,-1), 4),
        ("BOTTOMPADDING",(0,0),(-1,-1), 4),
        ("LEFTPADDING",  (0,0),(-1,-1), 6),
        ("RIGHTPADDING", (0,0),(-1,-1), 6),
        ("LINEBELOW",    (0,0),(-1,-2), 0.4, thin),
        ("LINEAFTER",    (1,0),(1,-1),  0.6, thin),   # divider between 2 col groups
        ("ROWBACKGROUNDS",(0,0),(-1,-1),[C_WHITE, C_GRAY_LIGHT]),
    ]
    field_tbl.setStyle(TableStyle(field_tbl_style))
    elems.append(field_tbl)
    elems.append(Spacer(1, 4*mm))

    # ── Classification history table ─────────────────────
    etype_list = rec.get("enterprise_type", [])

    subhdr_s = ParagraphStyle("_sh", fontName="Helvetica-Bold", fontSize=8.5,
                               textColor=C_BLUE, spaceBefore=2, spaceAfter=2)
    elems.append(Paragraph("Enterprise Type — Classification History", subhdr_s))

    if etype_list:
        hdrs    = ["Classification Year", "Enterprise Type", "Classification Date"]
        hdr_row_data = [Paragraph(h, S["tbl_hdr"]) for h in hdrs]
        tbl_data     = [hdr_row_data]

        for entry in etype_list:
            yr   = entry.get("classification_year", "") or "—"
            etype= entry.get("enterprise_type",     "") or "—"
            dt   = entry.get("classification_date", "") or "—"

            tc, bg = TYPE_COLOR.get(etype, (C_GRAY, C_GRAY_LIGHT))
            etype_s = ParagraphStyle(f"_et_{yr}", fontName="Helvetica-Bold",
                                      fontSize=8, textColor=tc, alignment=TA_CENTER)
            tbl_data.append([
                Paragraph(yr,    S["tbl_cell"]),
                Paragraph(etype, etype_s),
                Paragraph(dt,    S["tbl_cell"]),
            ])

        ct = Table(tbl_data, colWidths=[50*mm, 50*mm, 60*mm])
        ct_style = [
            ("BACKGROUND",    (0,0),(-1,0),  C_BLUE),
            ("ROWBACKGROUNDS",(0,1),(-1,-1), [C_WHITE, C_BLUE_LIGHT]),
            ("GRID",          (0,0),(-1,-1), 0.4, C_GRAY_LINE),
            ("VALIGN",        (0,0),(-1,-1), "MIDDLE"),
            ("TOPPADDING",    (0,0),(-1,-1), 5),
            ("BOTTOMPADDING", (0,0),(-1,-1), 5),
        ]
        # Per-row background for enterprise type column
        for ri, entry in enumerate(etype_list, 1):
            etype = entry.get("enterprise_type","")
            _, bg = TYPE_COLOR.get(etype, (C_GRAY, C_GRAY_LIGHT))
            ct_style.append(("BACKGROUND", (1,ri),(1,ri), bg))

        ct.setStyle(TableStyle(ct_style))
        elems.append(ct)
    else:
        elems.append(Paragraph(
            "No classification history found on portal for this record.",
            S["note"]))

    elems.append(Spacer(1, 6*mm))
    elems.append(HRFlowable(width="100%", thickness=0.5, color=C_GRAY_LINE))
    elems.append(Spacer(1, 4*mm))

    return [KeepTogether(elems[:4])] + elems[4:]   # keep header + id bar together


# ── Classification summary page ───────────────────────────

def classification_summary(S, records):
    elems = [PageBreak()]
    elems.append(Paragraph("Classification Summary", S["section_hdr"]))

    verified = [r for r in records if r.get("vstatus") == "Verified"]

    # Collect all unique years across all records
    all_years = sorted({
        e.get("classification_year","")
        for r in verified
        for e in r.get("enterprise_type", [])
        if e.get("classification_year","")
    })

    if not all_years:
        elems.append(Paragraph("No classification data available.", S["note"]))
        return elems

    elems.append(Paragraph(
        "Enterprise type per business per classification year. "
        "Colors: blue = Micro, green = Small, yellow = Medium.",
        S["note"]))
    elems.append(Spacer(1, 3*mm))

    # Build year→type lookup per record
    def get_type(rec, yr):
        for e in rec.get("enterprise_type",[]):
            if e.get("classification_year","") == yr:
                return e.get("enterprise_type","")
        return ""

    # Table: rows = enterprises, cols = years
    name_w = 55*mm
    yr_w   = min(14*mm, (W - 40*mm - name_w) / max(len(all_years),1))

    hdr = [Paragraph("Enterprise", S["tbl_hdr"])] + \
          [Paragraph(yr, S["tbl_hdr"]) for yr in all_years]
    tbl_data = [hdr]

    for rec in verified:
        name = (rec.get("name_of_enterprise","") or rec.get("udyam_id",""))[:40]
        row  = [Paragraph(name, S["tbl_cell_l"])]
        for yr in all_years:
            etype = get_type(rec, yr)
            tc, _ = TYPE_COLOR.get(etype,(C_GRAY, C_GRAY_LIGHT))
            es = ParagraphStyle(f"_s{yr}{name[:3]}", fontName="Helvetica-Bold",
                                 fontSize=7, textColor=tc, alignment=TA_CENTER)
            row.append(Paragraph(etype or "—", es))
        tbl_data.append(row)

    col_widths = [name_w] + [yr_w]*len(all_years)
    tbl = Table(tbl_data, colWidths=col_widths, repeatRows=1)

    tbl_cmds = [
        ("BACKGROUND",    (0,0),(-1,0),  C_BLUE),
        ("ROWBACKGROUNDS",(0,1),(-1,-1), [C_WHITE, C_GRAY_LIGHT]),
        ("GRID",          (0,0),(-1,-1), 0.4, C_GRAY_LINE),
        ("VALIGN",        (0,0),(-1,-1), "MIDDLE"),
        ("TOPPADDING",    (0,0),(-1,-1), 4),
        ("BOTTOMPADDING", (0,0),(-1,-1), 4),
        ("LEFTPADDING",   (0,0),(0,-1),  6),
    ]
    for ri, rec in enumerate(verified, 1):
        for ci, yr in enumerate(all_years, 1):
            etype = get_type(rec, yr)
            _, bg = TYPE_COLOR.get(etype, (C_GRAY, C_GRAY_LIGHT))
            if etype:
                tbl_cmds.append(("BACKGROUND",(ci,ri),(ci,ri), bg))

    tbl.setStyle(TableStyle(tbl_cmds))
    elems.append(tbl)
    elems.append(Spacer(1, 8*mm))

    # Count per year
    elems.append(Paragraph("Type Count per Year", S["section_hdr"]))
    count_hdrs = ["Year","Micro","Small","Medium","Total Verified"]
    count_data = [[Paragraph(h, S["tbl_hdr"]) for h in count_hdrs]]
    for yr in all_years:
        micro  = sum(1 for r in verified if get_type(r,yr)=="Micro")
        small  = sum(1 for r in verified if get_type(r,yr)=="Small")
        medium = sum(1 for r in verified if get_type(r,yr)=="Medium")
        tot    = micro+small+medium
        count_data.append([Paragraph(x, S["tbl_cell"]) for x in
                           [yr, str(micro), str(small), str(medium), str(tot)]])

    ct = Table(count_data, colWidths=[35*mm,30*mm,30*mm,30*mm,35*mm])
    ct.setStyle(TableStyle([
        ("BACKGROUND",    (0,0),(-1,0),  C_BLUE),
        ("ROWBACKGROUNDS",(0,1),(-1,-1), [C_WHITE, C_BLUE_LIGHT]),
        ("GRID",          (0,0),(-1,-1), 0.4, C_GRAY_LINE),
        ("VALIGN",        (0,0),(-1,-1), "MIDDLE"),
        ("TOPPADDING",    (0,0),(-1,-1), 5),
        ("BOTTOMPADDING", (0,0),(-1,-1), 5),
    ]))
    elems.append(ct)
    return elems


# ── Excel generation ──────────────────────────────────────

def generate_xlsx(records, xlsx_path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("⚠️  openpyxl not installed — skipping Excel. Run: pip install openpyxl")
        return

    wb   = Workbook()
    thin = Side(style="thin", color="E5E7EB")
    bdr  = Border(left=thin, right=thin, top=thin, bottom=thin)

    def hdr_cell(ws, row, col, text, bg="1A56DB"):
        c = ws.cell(row=row, column=col, value=text)
        c.font      = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        c.fill      = PatternFill("solid", start_color=bg)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border    = bdr
        return c

    def data_cell(ws, row, col, text, bg="FFFFFF", bold=False, center=False):
        c = ws.cell(row=row, column=col, value=text)
        c.font      = Font(name="Arial", size=9, bold=bold)
        c.fill      = PatternFill("solid", start_color=bg)
        c.alignment = Alignment(vertical="center",
                                horizontal="center" if center else "left",
                                wrap_text=True)
        c.border    = bdr
        return c

    STATUS_BG = {
        "Verified": "D1FAE5", "Invalid": "FEE2E2",
        "Invalid Format": "FEE2E2", "Error": "FEF3C7", "Timeout": "FEF3C7",
    }
    TYPE_BG = {"Micro": "DBEAFE", "Small": "D1FAE5", "Medium": "FEF3C7"}

    # ── Sheet 1: Enterprise Details ───────────────────────
    ws1 = wb.active
    ws1.title = "Enterprise Details"

    # Title
    ws1.merge_cells("A1:N1")
    t = ws1["A1"]
    t.value     = "Udyam Registration — Enterprise Details"
    t.font      = Font(name="Arial", bold=True, size=13, color="FFFFFF")
    t.fill      = PatternFill("solid", start_color="1A56DB")
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 26

    ws1.merge_cells("A2:N2")
    s = ws1["A2"]
    s.value     = f"Generated: {datetime.now().strftime('%d %B %Y, %H:%M')}   |   Total Records: {len(records)}"
    s.font      = Font(name="Arial", size=9, italic=True, color="6B7280")
    s.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 18

    cols1 = [
        ("Udyam ID",              "udyam_id",              24),
        ("Name of Enterprise",    "name_of_enterprise",    32),
        ("Date of Incorporation", "date_of_incorporation", 20),
        ("Major Activity",        "major_activity",        18),
        ("Social Category",       "social_category",       16),
        ("Date of Commencement",  "date_of_commencement",  20),
        ("Org Type",              "org_type",              20),
        ("NIC Code",              "nic_code",              12),
        ("State",                 "state",                 16),
        ("District",              "district",              16),
        ("DIC",                   "dic",                   16),
        ("MSME-DFO",              "msme_dfo",              18),
        ("Date of Udyam Reg",     "date_of_udyam_reg",     20),
        ("Verification",          "vstatus",               14),
    ]

    for ci, (hdr, _, w) in enumerate(cols1, 1):
        hdr_cell(ws1, 3, ci, hdr)
        ws1.column_dimensions[get_column_letter(ci)].width = w
    ws1.row_dimensions[3].height = 24

    for ri, rec in enumerate(records, 4):
        bg  = "F9FAFB" if ri % 2 == 0 else "FFFFFF"
        vs  = rec.get("vstatus", "")
        for ci, (_, key, _) in enumerate(cols1, 1):
            val  = rec.get(key, "") or ""
            cbg  = STATUS_BG.get(vs, bg) if key == "vstatus" else bg
            bold = key == "vstatus"
            data_cell(ws1, ri, ci, val, cbg, bold=bold,
                      center=(key in ("vstatus","nic_code","date_of_incorporation",
                                      "date_of_commencement","date_of_udyam_reg")))
        ws1.row_dimensions[ri].height = 20

    ws1.freeze_panes = "A4"

    # ── Sheet 2: Classification History ───────────────────
    ws2 = wb.create_sheet("Classification History")

    ws2.merge_cells("A1:F1")
    t2 = ws2["A1"]
    t2.value     = "Udyam Registration — Enterprise Type Classification History"
    t2.font      = Font(name="Arial", bold=True, size=13, color="FFFFFF")
    t2.fill      = PatternFill("solid", start_color="1A56DB")
    t2.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 26

    cols2 = [
        ("#",                    6),
        ("Udyam ID",            24),
        ("Name of Enterprise",  32),
        ("Classification Year", 20),
        ("Enterprise Type",     18),
        ("Classification Date", 20),
    ]
    for ci, (hdr, w) in enumerate(cols2, 1):
        hdr_cell(ws2, 2, ci, hdr)
        ws2.column_dimensions[get_column_letter(ci)].width = w
    ws2.row_dimensions[2].height = 22

    ri = 3
    for rec in records:
        udyam_id = rec.get("udyam_id", "")
        name     = rec.get("name_of_enterprise", "") or ""
        etype_list = rec.get("enterprise_type", [])

        if not etype_list:
            # Still add a row so enterprise is visible
            bg = "FEF3C7"
            data_cell(ws2, ri, 1, ri - 2,  bg, center=True)
            data_cell(ws2, ri, 2, udyam_id, bg)
            data_cell(ws2, ri, 3, name,     bg)
            data_cell(ws2, ri, 4, "—",      bg, center=True)
            data_cell(ws2, ri, 5, "No data", bg, center=True)
            data_cell(ws2, ri, 6, "—",      bg, center=True)
            ws2.row_dimensions[ri].height = 20
            ri += 1
            continue

        for ei, entry in enumerate(etype_list):
            yr    = entry.get("classification_year", "") or "—"
            etype = entry.get("enterprise_type",     "") or "—"
            dt    = entry.get("classification_date", "") or "—"
            bg    = TYPE_BG.get(etype, "FFFFFF")
            row_bg = "F9FAFB" if ri % 2 == 0 else "FFFFFF"

            # Show enterprise info only on first row, merge visually with shading
            data_cell(ws2, ri, 1, ri - 2 if ei == 0 else "", row_bg, center=True)
            data_cell(ws2, ri, 2, udyam_id if ei == 0 else "", row_bg)
            data_cell(ws2, ri, 3, name     if ei == 0 else "", row_bg)
            data_cell(ws2, ri, 4, yr,    bg, center=True)
            data_cell(ws2, ri, 5, etype, bg, bold=True, center=True)
            data_cell(ws2, ri, 6, dt,    bg, center=True)
            ws2.row_dimensions[ri].height = 20
            ri += 1

    ws2.freeze_panes = "A3"

    wb.save(xlsx_path)
    print(f"💾 Excel saved → {xlsx_path}")


# ── Main generate ──────────────────────────────────────────

def generate_pdf(json_path, out_path=None):
    print(f"📂 Loading: {json_path}")
    with open(json_path, encoding="utf-8") as f:
        records = json.load(f)
    print(f"✅ Loaded {len(records)} records")

    base = os.path.splitext(json_path)[0]
    ts   = datetime.now().strftime("%Y%m%d_%H%M")

    if not out_path:
        out_path = f"{base}_report_{ts}.pdf"

    xlsx_path = f"{base}_report_{ts}.xlsx"

    # ── Generate Excel ────────────────────────────────────
    generate_xlsx(records, xlsx_path)

    # ── Generate PDF ──────────────────────────────────────
    S   = make_styles()
    doc = SimpleDocTemplate(
        out_path, pagesize=A4,
        rightMargin=20*mm, leftMargin=20*mm,
        topMargin=20*mm,   bottomMargin=14*mm,
        title="Udyam Verification Report",
    )

    story = []
    story += cover_page(S, records)

    story.append(Paragraph("Enterprise Details", S["section_hdr"]))
    story.append(Spacer(1, 2*mm))
    for i, rec in enumerate(records, 1):
        story += enterprise_block(S, rec, i)

    story += classification_summary(S, records)

    doc.build(story, canvasmaker=make_canvas_class(len(records)))

    size_kb = os.path.getsize(out_path) // 1024
    print(f"💾 PDF saved  → {out_path}  ({size_kb} KB)")
    return out_path, xlsx_path


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python udyam_pdf_report.py <verified.json> [output.pdf]")
        sys.exit(1)
    jp = sys.argv[1]
    op = sys.argv[2] if len(sys.argv) > 2 else None
    if not os.path.exists(jp):
        print(f"❌ File not found: {jp}")
        sys.exit(1)
    pdf_path, xlsx_path = generate_pdf(jp, op)
    print(f"\n✅ Done!")
    print(f"   PDF   → {pdf_path}")
    print(f"   Excel → {xlsx_path}")
